"""Read Excel Tables (ListObjects) and CSV into import payloads.

Sheets alone are not imported — only Insert→Table ranges (and CSV as one table).
"""

from __future__ import annotations

import csv
import io
from datetime import date, datetime
from typing import Any

from openpyxl.utils import range_boundaries


MAX_PREVIEW_ROWS = 5
MAX_IMPORT_ROWS = 5000
MAX_COLS = 80


def _looks_like_excel_date_format(number_format: str) -> bool:
    fmt = (number_format or "").strip().lower()
    if not fmt or fmt in {"general", "@", "0", "0.00"}:
        return False
    if "fa-ir" in fmt or "fa_ir" in fmt:
        return True
    return any(tok in fmt for tok in ("yy", "mm", "dd", "yyyy", "m/", "d/", "/m", "/d"))


def _cell_str(value: Any, *, number_format: str = "") -> str:
    """Normalize a worksheet cell to a transferable string.

    Date/datetime values and Excel serials with a date (incl. fa-IR) number format
    become Jalali ``YYYY/MM/DD`` (e.g. ``1405/06/01``) so the system stores/display
    شمسی — never raw میلادی years in the Excel grid.
    """
    from catalog.jalali_dates import format_jalali_slash, excel_serial_to_gregorian

    if value is None:
        return ""
    if isinstance(value, datetime):
        return format_jalali_slash(value.date())
    if isinstance(value, date):
        return format_jalali_slash(value)
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if _looks_like_excel_date_format(number_format):
            as_date = excel_serial_to_gregorian(value)
            if as_date is not None:
                return format_jalali_slash(as_date)
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return str(value).strip()
    return str(value).strip()


def _normalize_headers(raw: list[Any], width: int) -> list[str]:
    headers: list[str] = []
    seen: dict[str, int] = {}
    for i in range(width):
        base = _cell_str(raw[i]) if i < len(raw) else ""
        if not base:
            base = f"ستون {i + 1}"
        base = base[:120]
        n = seen.get(base, 0)
        seen[base] = n + 1
        headers.append(base if n == 0 else f"{base} ({n + 1})")
    return headers


def _split_header_rows(matrix: list[list[Any]]) -> tuple[list[str], list[list[str]]]:
    if not matrix:
        return [], []
    width = min(MAX_COLS, max((len(r) for r in matrix), default=0))
    if width <= 0:
        return [], []
    headers = _normalize_headers(matrix[0], width)
    rows: list[list[str]] = []
    for raw in matrix[1 : MAX_IMPORT_ROWS + 1]:
        row = [_cell_str(raw[i]) if i < len(raw) else "" for i in range(width)]
        if any(row):
            rows.append(row)
    return headers, rows


def _matrix_from_table(ws, table) -> tuple[list[str], list[list[str]]]:
    min_col, min_row, max_col, max_row = range_boundaries(table.ref)
    width = min(MAX_COLS, max_col - min_col + 1)
    if width <= 0:
        return [], []
    header_count = int(getattr(table, "headerRowCount", None) or 1)
    header_count = max(1, min(header_count, max_row - min_row + 1))

    # Prefer structured table column names when present
    col_names = []
    try:
        col_names = [c.name for c in (table.tableColumns or [])]
    except Exception:  # noqa: BLE001
        col_names = []
    if col_names and len(col_names) >= width:
        headers = _normalize_headers(col_names[:width], width)
    else:
        header_cells = [
            ws.cell(row=min_row, column=min_col + i).value for i in range(width)
        ]
        headers = _normalize_headers(header_cells, width)

    data_start = min_row + header_count
    rows: list[list[str]] = []
    for r in range(data_start, max_row + 1):
        if len(rows) >= MAX_IMPORT_ROWS:
            break
        # Skip totals row if marked
        totals = int(getattr(table, "totalsRowCount", None) or 0)
        if totals and r > max_row - totals:
            continue
        row = []
        for i in range(width):
            cell = ws.cell(row=r, column=min_col + i)
            row.append(_cell_str(cell.value, number_format=str(cell.number_format or "")))
        if any(row):
            rows.append(row)
    return headers, rows


def _table_id(sheet_name: str, display_name: str) -> str:
    return f"{sheet_name}::{display_name}"


def preview_workbook(uploaded_file) -> list[dict]:
    """Return Excel Table previews (not whole sheets).

    Each item: name (table display name), sheet_name, table_id, headers, counts.
    """
    name = (getattr(uploaded_file, "name", "") or "").lower()
    content = uploaded_file.read()
    if hasattr(uploaded_file, "seek"):
        uploaded_file.seek(0)

    if name.endswith(".csv"):
        text = content.decode("utf-8-sig", errors="replace")
        reader = csv.reader(io.StringIO(text))
        matrix = [list(r) for r in reader]
        headers, rows = _split_header_rows(matrix)
        table_name = (getattr(uploaded_file, "name", "") or "Table1").rsplit("/", 1)[-1]
        if table_name.lower().endswith(".csv"):
            table_name = table_name[:-4] or "Table1"
        return [{
            "name": table_name[:200] or "Table1",
            "sheet_name": "CSV",
            "table_id": _table_id("CSV", table_name[:200] or "Table1"),
            "ref": "",
            "headers": headers,
            "row_count": len(rows),
            "column_count": len(headers),
            "preview_rows": rows[:MAX_PREVIEW_ROWS],
        }]

    from openpyxl import load_workbook

    # Tables are unavailable in read_only mode.
    wb = load_workbook(io.BytesIO(content), data_only=True)
    tables_out: list[dict] = []
    try:
        for ws in wb.worksheets:
            if not getattr(ws, "tables", None):
                continue
            for key in list(ws.tables.keys()):
                table = ws.tables[key]
                display = str(getattr(table, "displayName", None) or table.name or key)
                headers, rows = _matrix_from_table(ws, table)
                tables_out.append({
                    "name": display[:200],
                    "sheet_name": str(ws.title)[:200],
                    "table_id": _table_id(str(ws.title), display),
                    "ref": str(table.ref or ""),
                    "headers": headers,
                    "row_count": len(rows),
                    "column_count": len(headers),
                    "preview_rows": rows[:MAX_PREVIEW_ROWS],
                })
    finally:
        wb.close()
    return tables_out


def read_table_data(
    uploaded_file,
    *,
    sheet_name: str,
    table_name: str,
) -> tuple[list[str], list[list[str]]]:
    """Return headers+rows for one Excel Table by sheet + display name."""
    name = (getattr(uploaded_file, "name", "") or "").lower()
    content = uploaded_file.read()
    if hasattr(uploaded_file, "seek"):
        uploaded_file.seek(0)

    if name.endswith(".csv"):
        text = content.decode("utf-8-sig", errors="replace")
        reader = csv.reader(io.StringIO(text))
        matrix = [list(r) for r in reader]
        return _split_header_rows(matrix)

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content), data_only=True)
    try:
        target_ws = None
        for ws in wb.worksheets:
            if str(ws.title) == sheet_name:
                target_ws = ws
                break
        if target_ws is None:
            raise ValueError(f"شیت «{sheet_name}» در فایل یافت نشد.")
        if not getattr(target_ws, "tables", None):
            raise ValueError(f"در شیت «{sheet_name}» هیچ Table یافت نشد.")

        target = None
        for key in target_ws.tables:
            table = target_ws.tables[key]
            display = str(getattr(table, "displayName", None) or table.name or key)
            if display == table_name or str(key) == table_name:
                target = table
                break
        if target is None:
            raise ValueError(
                f"جدول «{table_name}» در شیت «{sheet_name}» یافت نشد."
            )
        return _matrix_from_table(target_ws, target)
    finally:
        wb.close()


# Backward-compatible alias used by older call sites / tests
def read_sheet_data(uploaded_file, sheet_name: str) -> tuple[list[str], list[list[str]]]:
    """Deprecated: whole-sheet import. Prefer read_table_data."""
    name = (getattr(uploaded_file, "name", "") or "").lower()
    content = uploaded_file.read()
    if hasattr(uploaded_file, "seek"):
        uploaded_file.seek(0)
    if name.endswith(".csv"):
        text = content.decode("utf-8-sig", errors="replace")
        reader = csv.reader(io.StringIO(text))
        return _split_header_rows([list(r) for r in reader])

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content), data_only=True)
    try:
        for ws in wb.worksheets:
            if str(ws.title) == sheet_name:
                # If sheet has tables, return the first table only
                if ws.tables:
                    key = next(iter(ws.tables.keys()))
                    return _matrix_from_table(ws, ws.tables[key])
                matrix = []
                for i, row in enumerate(ws.iter_rows(values_only=True)):
                    if i > MAX_IMPORT_ROWS:
                        break
                    matrix.append(list(row) if row is not None else [])
                return _split_header_rows(matrix)
        raise ValueError(f"شیت «{sheet_name}» در فایل یافت نشد.")
    finally:
        wb.close()
