"""Tests for Excel workbook import, grid edit, and report source wiring."""

from __future__ import annotations

import io
import json

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.test import TestCase
from django.urls import reverse

from catalog.models import ExcelTable, ExcelUpload
from reports.columns import get_column_groups, run_report

User = get_user_model()


def _make_xlsx_bytes() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws1 = wb.active
    ws1.title = "موجودی"
    ws1.append(["کد", "نام", "موجودی"])
    ws1.append(["A1", "قطعه یک", 10])
    ws1.append(["A2", "قطعه دو", 5])
    ws2 = wb.create_sheet("قیمت")
    ws2.append(["کد", "قیمت"])
    ws2.append(["A1", 1000])
    ws3 = wb.create_sheet("خالی")
    ws3.append(["ستون"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class ExcelManagementTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        call_command("seed_demo")
        cls.admin = User.objects.get(username="admin")
        cls.expert = User.objects.get(username="expert")
        cls.viewer = User.objects.get(username="viewer")

    def _xlsx(self, name="sample.xlsx"):
        return SimpleUploadedFile(
            name,
            _make_xlsx_bytes(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    def test_list_requires_login(self):
        resp = self.client.get(reverse("excel_list"))
        self.assertEqual(resp.status_code, 302)

    def test_viewer_can_list_but_not_import(self):
        self.client.login(username="viewer", password="erp12345")
        self.assertEqual(self.client.get(reverse("excel_list")).status_code, 200)
        self.assertEqual(self.client.get(reverse("excel_import")).status_code, 403)

    def test_preview_and_import_selected_sheets(self):
        self.client.login(username="expert", password="erp12345")
        preview = self.client.post(
            reverse("excel_preview"),
            {"file": self._xlsx("multi.xlsx")},
        )
        self.assertEqual(preview.status_code, 200)
        data = preview.json()
        self.assertTrue(data["ok"])
        self.assertEqual(data["filename"], "multi.xlsx")
        names = [s["name"] for s in data["sheets"]]
        self.assertEqual(names, ["موجودی", "قیمت", "خالی"])

        confirm = self.client.post(
            reverse("excel_import_confirm"),
            {
                "file": self._xlsx("multi.xlsx"),
                "title": "فایل تست موجودی",
                "selected_sheets": json.dumps([
                    {"sheet": "موجودی", "name": "جدول موجودی"},
                    {"sheet": "قیمت", "name": "قیمت"},
                ]),
            },
        )
        self.assertEqual(confirm.status_code, 200)
        payload = confirm.json()
        self.assertTrue(payload["ok"])
        self.assertEqual(payload["table_count"], 2)
        upload = ExcelUpload.objects.get(pk=payload["upload_id"])
        self.assertEqual(upload.title, "فایل تست موجودی")
        tables = list(upload.tables.order_by("order"))
        self.assertEqual([t.name for t in tables], ["جدول موجودی", "قیمت"])
        self.assertEqual([t.sheet_name for t in tables], ["موجودی", "قیمت"])
        self.assertEqual(tables[0].headers[0], "کد")
        self.assertEqual(tables[0].row_count, 2)
        # Must import the named sheet data, not silently fall back to sheet 1
        self.assertEqual(tables[1].headers, ["کد", "قیمت"])
        self.assertEqual(tables[1].rows[0][1], "1000")

        detail = self.client.get(reverse("excel_detail", args=[upload.pk]))
        self.assertEqual(detail.status_code, 200)
        self.assertContains(detail, "موجودی")
        self.assertContains(detail, "قیمت")
        self.assertContains(detail, "excel-grid")

    def test_save_table_headers_name_and_layout(self):
        self.client.login(username="expert", password="erp12345")
        upload = ExcelUpload.objects.create(title="ت", uploaded_by=self.expert)
        table = ExcelTable.objects.create(
            upload=upload,
            name="شیت۱",
            sheet_name="شیت۱",
            headers=["A", "B"],
            rows=[["1", "2"]],
        )
        resp = self.client.post(
            reverse("excel_table_save", args=[table.pk]),
            data=json.dumps({
                "name": "جدول ویرایش‌شده",
                "headers": ["A", "B", "C"],
                "rows": [["1", "2", "3"], ["4", "5", "6"]],
                "layout": {"colWidths": [100, 140, 80], "rowHeights": [30, 32]},
            }),
            content_type="application/json",
        )
        self.assertEqual(resp.status_code, 200)
        self.assertTrue(resp.json()["ok"])
        table.refresh_from_db()
        self.assertEqual(table.name, "جدول ویرایش‌شده")
        self.assertEqual(table.column_count, 3)
        self.assertEqual(table.row_count, 2)
        self.assertEqual(table.rows[1][2], "6")
        self.assertEqual(table.layout.get("colWidths"), [100, 140, 80])

    def test_missing_sheet_does_not_import_sheet_one(self):
        self.client.login(username="expert", password="erp12345")
        confirm = self.client.post(
            reverse("excel_import_confirm"),
            {
                "file": self._xlsx("multi.xlsx"),
                "title": "بد",
                "selected_sheets": json.dumps([{"sheet": "وجودندارد", "name": "X"}]),
            },
        )
        self.assertEqual(confirm.status_code, 400)
        self.assertFalse(confirm.json()["ok"])
        self.assertFalse(ExcelUpload.objects.filter(title="بد").exists())

    def test_delete_table_manager_only(self):
        upload = ExcelUpload.objects.create(title="حذف", uploaded_by=self.admin)
        table = ExcelTable.objects.create(
            upload=upload, name="T1", sheet_name="T1", headers=["X"], rows=[["1"]],
        )
        self.client.login(username="expert", password="erp12345")
        forbidden = self.client.post(reverse("excel_table_delete", args=[table.pk]))
        self.assertEqual(forbidden.status_code, 403)
        self.assertTrue(ExcelTable.objects.filter(pk=table.pk).exists())

        self.client.login(username="admin", password="erp12345")
        ok = self.client.post(reverse("excel_table_delete", args=[table.pk]))
        self.assertEqual(ok.status_code, 302)
        self.assertFalse(ExcelTable.objects.filter(pk=table.pk).exists())

    def test_excel_tables_appear_in_report_sources(self):
        upload = ExcelUpload.objects.create(title="منبع", uploaded_by=self.admin)
        table = ExcelTable.objects.create(
            upload=upload,
            name="جدول فروش",
            sheet_name="فروش",
            headers=["کد", "مقدار"],
            rows=[["C1", "9"], ["C2", "3"]],
        )
        groups = get_column_groups()
        ids = [g["id"] for g in groups]
        self.assertIn(table.source_id, ids)
        group = next(g for g in groups if g["id"] == table.source_id)
        self.assertEqual(group["label"], "جدول فروش")
        labels = [c[1] for c in group["columns"]]
        self.assertEqual(labels, ["کد", "مقدار"])

        headers, rows, _payloads, _deeper = run_report(
            table.source_id,
            [
                {"key": "col_0", "source": table.source_id, "level": 1, "label": "کد"},
                {"key": "col_1", "source": table.source_id, "level": 1, "label": "مقدار"},
            ],
        )
        self.assertEqual(headers, ["کد", "مقدار"])
        self.assertEqual(rows, [["C1", "9"], ["C2", "3"]])
