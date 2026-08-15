"""Import printable A4 forms from an Excel workbook into PrintForm frames."""

from __future__ import annotations

import uuid
from io import BytesIO
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .models import PrintForm

A4_WIDTH_MM = 210
A4_HEIGHT_MM = 297


def _next_form_number(owner) -> int:
    used = set(PrintForm.objects.filter(owner=owner).values_list("number", flat=True))
    for n in range(1, 1000):
        if n not in used:
            return n
    raise ValueError("شماره فرم آزاد باقی نمانده است.")


def _cell_text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _sheet_to_frames(ws, page_w: float = A4_WIDTH_MM, page_h: float = A4_HEIGHT_MM) -> list[dict]:
    """Map non-empty cells (and merges) into designer frames in millimetres."""
    max_row = min(ws.max_row or 1, 80)
    max_col = min(ws.max_column or 1, 20)
    if max_row < 1 or max_col < 1:
        return []

    col_w = page_w / max_col
    row_h = page_h / max_row

    merged_map: dict[tuple[int, int], tuple[int, int, int, int]] = {}
    for rng in ws.merged_cells.ranges:
        min_r, min_c, max_r, max_c = rng.min_row, rng.min_col, rng.max_row, rng.max_col
        for r in range(min_r, max_r + 1):
            for c in range(min_c, max_c + 1):
                merged_map[(r, c)] = (min_r, min_c, max_r, max_c)

    frames: list[dict[str, Any]] = []
    seen_merges: set[tuple[int, int, int, int]] = set()

    for r in range(1, max_row + 1):
        for c in range(1, max_col + 1):
            cell = ws.cell(r, c)
            text = _cell_text(cell.value)
            merge = merged_map.get((r, c))
            if merge:
                if merge in seen_merges:
                    continue
                if (r, c) != (merge[0], merge[1]):
                    continue
                seen_merges.add(merge)
                min_r, min_c, max_r, max_c = merge
                text = _cell_text(ws.cell(min_r, min_c).value)
                x = (min_c - 1) * col_w
                y = (min_r - 1) * row_h
                w = (max_c - min_c + 1) * col_w
                h = (max_r - min_r + 1) * row_h
            else:
                if not text:
                    continue
                x = (c - 1) * col_w
                y = (r - 1) * row_h
                w = col_w
                h = row_h

            kind = "header" if r <= 2 and len(text) > 0 else "field"
            if merge and not text:
                kind = "box"
            frames.append(
                {
                    "id": uuid.uuid4().hex[:10],
                    "kind": kind,
                    "label": text or get_column_letter(c) + str(r),
                    "x": round(x, 2),
                    "y": round(y, 2),
                    "width": round(max(w, 8), 2),
                    "height": round(max(h, 6), 2),
                    "font_size": 10 if kind != "header" else 12,
                    "align": "center",
                }
            )
    return frames


def import_forms_from_excel(file_obj, owner, created_by=None) -> list[PrintForm]:
    """Create one PrintForm per worksheet (A4). Returns created forms."""
    data = file_obj.read() if hasattr(file_obj, "read") else file_obj
    wb = load_workbook(BytesIO(data), data_only=True)
    created: list[PrintForm] = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        frames = _sheet_to_frames(ws)
        if not frames:
            continue
        title = (sheet_name or "فرم واردشده").strip()[:200] or "فرم واردشده"
        number = _next_form_number(owner)
        obj = PrintForm.objects.create(
            owner=owner,
            created_by=created_by or owner,
            title=title,
            description="واردشده از اکسل",
            number=number,
            frames=frames,
            page_width_mm=A4_WIDTH_MM,
            page_height_mm=A4_HEIGHT_MM,
            is_standard=False,
        )
        created.append(obj)
    if not created:
        raise ValueError("در فایل اکسل محتوای قابل تبدیل به فرم یافت نشد.")
    return created
